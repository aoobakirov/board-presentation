# -*- coding: utf-8 -*-
"""Design toolkit replicating the KazBioFert board-deck corporate style (720x405 pt)."""
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from copy import deepcopy
import re

# ---- palette (extracted from deck) ----
NAVY   = RGBColor(0x13,0x29,0x4B)
NAVYDK = RGBColor(0x0C,0x1D,0x38)
STEEL  = RGBColor(0x2F,0x66,0x90)
COPPER = RGBColor(0xC4,0x77,0x2F)
TAN    = RGBColor(0xE0,0xA4,0x6A)
INK    = RGBColor(0x2A,0x37,0x46)
GRAY   = RGBColor(0x6B,0x7A,0x8D)
LBLUE  = RGBColor(0xAE,0xC8,0xDE)
WHITE  = RGBColor(0xFF,0xFF,0xFF)
BORDER = RGBColor(0xDC,0xE3,0xEB)
PANEL  = RGBColor(0xF4,0xF6,0xF9)
MIST   = RGBColor(0xEA,0xEF,0xF4)
GREENT = RGBColor(0x5B,0x8C,0x6E)   # subtle positive accent

EMU_PT = 12700

def _set_bg(slide, color):
    # remove existing bg then set solid
    cs = slide._element.spPr if hasattr(slide._element,'spPr') else None
    bgPr = slide._element.find(qn('p:cSld')+'/'+qn('p:bg'))
    csld = slide._element.find(qn('p:cSld'))
    old = csld.find(qn('p:bg'))
    if old is not None:
        csld.remove(old)
    bg = csld.makeelement(qn('p:bg'), {})
    bgPr = bg.makeelement(qn('p:bgPr'), {})
    fill = bgPr.makeelement(qn('a:solidFill'), {})
    clr = fill.makeelement(qn('a:srgbClr'), {'val': '%02X%02X%02X'%(color[0],color[1],color[2])})
    fill.append(clr); bgPr.append(fill)
    bgPr.append(bgPr.makeelement(qn('a:effectLst'), {}))
    bg.append(bgPr)
    csld.insert(0, bg)

def soft_shadow(shape, blur=101600, dist=25400, alpha=10000, dirv=5400000):
    spPr = shape._element.spPr
    # remove existing effectLst
    for e in spPr.findall(qn('a:effectLst')):
        spPr.remove(e)
    eff = spPr.makeelement(qn('a:effectLst'), {})
    sh = eff.makeelement(qn('a:outerShdw'), {
        'blurRad':str(blur),'dist':str(dist),'dir':str(dirv),'algn':'bl','rotWithShape':'0'})
    clr = sh.makeelement(qn('a:srgbClr'), {'val':'000000'})
    a = clr.makeelement(qn('a:alpha'), {'val':str(alpha)})
    clr.append(a); sh.append(clr); eff.append(sh); spPr.append(eff)

def no_line(shape):
    shape.line.fill.background()

def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, round_=None, shadow=False):
    shp = MSO_SHAPE.ROUNDED_RECTANGLE if round_ is not None else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shp, Pt(x), Pt(y), Pt(w), Pt(h))
    if round_ is not None:
        try: s.adjustments[0] = round_
        except Exception: pass
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    s.shadow.inherit = False
    if shadow:
        soft_shadow(s)
    return s

def line_seg(slide, x, y, w, h, color, weight=1.0):
    from pptx.enum.shapes import MSO_CONNECTOR
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x), Pt(y), Pt(x+w), Pt(y+h))
    c.line.color.rgb = color; c.line.width = Pt(weight)
    c.shadow.inherit = False
    return c

def _apply_runs(tf, lines, align, anchor, space_after=None, line_spacing=None):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ('left','right','top','bottom'):
        setattr(tf, 'margin_'+m, 0)
    first = True
    for line in lines:  # line = list of run-dicts
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if space_after is not None: p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        if line_spacing is not None: p.line_spacing = line_spacing
        for rd in line:
            r = p.add_run(); r.text = rd['t']
            f = r.font
            f.size = Pt(rd.get('sz',10)); f.bold = rd.get('b',False)
            f.name = rd.get('f','Calibri')
            f.color.rgb = rd.get('c',INK)
            if rd.get('spc') is not None:
                r.font._rPr.set('spc', str(int(rd['spc']*100)))
            if rd.get('i'): f.italic = True

def text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=None, line_spacing=None):
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    _apply_runs(tb.text_frame, lines, align, anchor, space_after, line_spacing)
    return tb

def simple(slide, x, y, w, h, s, sz, b, c, f='Calibri', align=PP_ALIGN.LEFT,
           anchor=MSO_ANCHOR.TOP, spc=None, line_spacing=None):
    return text(slide, x, y, w, h, [[{'t':s,'sz':sz,'b':b,'c':c,'f':f,'spc':spc}]],
                align, anchor, line_spacing=line_spacing)

# ---- deck chrome ----
def eyebrow(slide, s, color=COPPER, x=40, y=24, w=641):
    return simple(slide, x, y, w, 16, s, 11, True, color, 'Calibri', spc=1.4)

def title(slide, s, color=NAVY, x=40, y=46, w=648, sz=21):
    return simple(slide, x, y, w, 28, s, sz, True, color, 'Arial')

def subtitle(slide, s, color=RGBColor(0x53,0x62,0x72), x=40, y=74, w=648, sz=10.5):
    return simple(slide, x, y, w, 18, s, sz, False, color, 'Calibri')

def footer(slide, page, section, brand_c=STEEL, sec_c=GRAY, on_dark=False):
    bc = LBLUE if on_dark else STEEL
    gc = RGBColor(0x8A,0x9A,0xAD) if on_dark else GRAY
    line_seg(slide, 40, 374, 641, 0, BORDER if not on_dark else RGBColor(0x2A,0x3D,0x5C), 0.75)
    text(slide, 40, 380, 432, 18, [[
        {'t':'KazBioFert / BBFF','sz':8.5,'b':True,'c':bc,'f':'Calibri'},
        {'t':'  ·  Совет директоров','sz':8.5,'b':False,'c':gc,'f':'Calibri'}]])
    simple(slide, 300, 380, 350, 18, section, 8, False, gc, 'Calibri', PP_ALIGN.RIGHT, spc=0.8)
    simple(slide, 652, 380, 29, 18, str(page), 9, True, gc, 'Calibri', PP_ALIGN.RIGHT)

def icon_chip(slide, x, y, size, fill=STEEL, glyph=None, gcolor=WHITE, gsize=None):
    rect(slide, x, y, size, size, fill=fill, round_=0.22)
    if glyph:
        simple(slide, x, y-1, size, size, glyph, gsize or size*0.5, True, gcolor, 'Calibri',
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

def kpi_chip(slide, x, y, w, h, big, unit, label, big_c=STEEL, on_dark=False, big_sz=23):
    card_fill = NAVYDK if on_dark else WHITE
    r = rect(slide, x, y, w, h, fill=card_fill, line=(RGBColor(0x24,0x3A,0x59) if on_dark else BORDER),
             line_w=1.0, round_=0.10, shadow=not on_dark)
    rect(slide, x+2, y+5, 3.0, h-10, fill=big_c)  # accent bar (inset)
    text(slide, x+14, y+6, w-22, 24, [[
        {'t':big,'sz':big_sz,'b':True,'c':big_c,'f':'Arial'},
        {'t':(' '+unit if unit else ''),'sz':10.5,'b':True,'c':(LBLUE if on_dark else GRAY),'f':'Calibri'}]],
        anchor=MSO_ANCHOR.BOTTOM)
    simple(slide, x+14, y+30, w-24, h-32, label, 8.6, False,
           (LBLUE if on_dark else INK), 'Calibri', line_spacing=1.0)
    return r

def source_note(slide, s, x=40, y=360, on_dark=False):
    c = RGBColor(0x6F,0x82,0x98) if on_dark else RGBColor(0x9A,0xA6,0xB3)
    return simple(slide, x, y, 632, 12, s, 7, False, c, 'Calibri')

def hrule(slide, y, x=40, w=641, color=BORDER, wt=0.75):
    return line_seg(slide, x, y, w, 0, color, wt)

def header_mark(slide, y=25, x=40, h=32, color=COPPER):
    """Thin vertical accent bar as a section marker; content indents to x+12."""
    return rect(slide, x, y, 3, h, fill=color)

def card_head(slide, x, y, w, s, color=NAVY, sz=9.5):
    return simple(slide, x, y, w, 14, s, sz, True, color, 'Calibri')

def clean_table(tbl):
    """Strip default banded table style; use plain no-grid so manual fills show cleanly."""
    tblPr = tbl._tbl.find(qn('a:tblPr'))
    if tblPr is None:
        tblPr = tbl._tbl.makeelement(qn('a:tblPr'), {}); tbl._tbl.insert(0, tblPr)
    tblPr.set('firstRow','0'); tblPr.set('bandRow','0')
    for sid in tblPr.findall(qn('a:tableStyleId')):
        tblPr.remove(sid)
    sid = tblPr.makeelement(qn('a:tableStyleId'), {})
    sid.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'  # No Style, No Grid
    tblPr.append(sid)

def cell_border(cell, sides=('bottom',), color=BORDER, w=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    tagmap = {'left':'a:lnL','right':'a:lnR','top':'a:lnT','bottom':'a:lnB'}
    for side in sides:
        tag = tagmap[side]
        for e in tcPr.findall(qn(tag)): tcPr.remove(e)
        ln = tcPr.makeelement(qn(tag), {'w':str(int(w*EMU_PT)),'cap':'flat'})
        sf = ln.makeelement(qn('a:solidFill'), {})
        c = sf.makeelement(qn('a:srgbClr'), {'val':'%02X%02X%02X'%(color[0],color[1],color[2])})
        sf.append(c); ln.append(sf)
        tcPr.append(ln)

def move_slide(prs, from_idx, to_idx):
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    el = ids[from_idx]
    sldIdLst.remove(el)
    sldIdLst.insert(to_idx, el)

def add_blank(prs, bg=WHITE):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    # strip any inherited placeholders
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    _set_bg(s, bg)
    return s
