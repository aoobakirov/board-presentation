# -*- coding: utf-8 -*-
"""Patch the user's manually-edited deck: migrate slides 23-24 (budget) to a
consistent 6-month / 1-е полугодие 2026 basis to match the ДДС table (266,5 млн ₸).
Operates on the user's file directly so all manual edits (ДДС table, tax note,
'Экономия 150 млн', RU wording, $450/75k market figures) are preserved."""
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData

SRC = "source/base_manual_v1.pptx"
OUT = "KazBioFert_Board_Presentation_FINAL_20260709.pptx"

prs = Presentation(SRC)

def set_full_text(shape, new):
    """Replace a shape's whole text, keeping the first run's formatting."""
    tf = shape.text_frame
    p = tf.paragraphs[0]
    if p.runs:
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.add_run().text = new
    # clear any extra paragraphs
    for extra in tf.paragraphs[1:]:
        for r in extra.runs:
            r.text = ''

# ---------- SLIDE 23 (index 22): execution chart -> 6 months ----------
s = prs.slides[22]
for sh in s.shapes:
    if sh.has_chart:
        cd = CategoryChartData()
        cd.categories = ['Операц.\n(выбытие)', 'Инвест.\n(выбытие)', 'Финанс.\n(поступл.)']
        cd.add_series('План', (124.6, 513.7, 637.8))
        cd.add_series('Факт', (76.8, 183.4, 266.5))
        sh.chart.replace_data(cd)
        # re-apply original series colours (План=light blue, Факт=steel)
        sh.chart.series[0].format.fill.solid()
        sh.chart.series[0].format.fill.fore_color.rgb = RGBColor(0xAE,0xC8,0xDE)
        sh.chart.series[1].format.fill.solid()
        sh.chart.series[1].format.fill.fore_color.rgb = RGBColor(0x2F,0x66,0x90)

txt_map = {
    'ИСПОЛНЕНИЕ БЮДЖЕТА ЗА 5 МЕСЯЦЕВ 2026': 'ИСПОЛНЕНИЕ БЮДЖЕТА ЗА 1 ПОЛУГОДИЕ 2026',
    'млн ₸ · план-факт за 5 месяцев 2026': 'млн ₸ · план-факт за 6 месяцев 2026',
    '+46,9': '+47,8',
    '+244,9': '+330,3',
    '−323,1': '−371,3',
    'БЮДЖЕТ 5 МЕС.': 'БЮДЖЕТ 6 МЕС.',
}
for sh in s.shapes:
    if sh.has_text_frame:
        cur = sh.text_frame.text.strip()
        if cur in txt_map:
            set_full_text(sh, txt_map[cur])

# ---------- SLIDE 24 (index 23): subtitle + remove orphan ----------
s = prs.slides[23]
to_remove = []
for sh in s.shapes:
    if sh.has_text_frame:
        cur = sh.text_frame.text.strip()
        if cur.startswith('Годовой бюджет'):
            set_full_text(sh, 'Годовой бюджет ~728 млн ₸; исполнено 266,5 млн ₸ за 1 полугодие (37%), '
                              '~462 млн ₸ — во II полугодии')
        elif cur == 'Итоговый годовой план, млн ₸':
            to_remove.append(sh)
for sh in to_remove:
    sh._element.getparent().remove(sh._element)

# ---------- COST OF GOODS -> $330/t (slides 15, 18, 19) ----------
# price kept at $450/t  =>  unit gross margin = 450-330 = $120/t = ~27%
def find_text(slide, exact):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == exact:
            return sh
    return None

yB, scale = 270.0, (270.0 - 150.0) / 630.0   # cost-bridge baseline & px/$ scale

# --- slide 18: pricing / economics ---
s = prs.slides[17]
set_full_text(find_text(s, 'Себестоимость ОМУ $233/т при цене реализации $450/т'),
              'Себестоимость ОМУ $330/т при цене реализации $450/т')
set_full_text(find_text(s, 'ОМУ NPK 8-21: сырьё $194 + производство (с пост. расходами) $39 = $233/т; валовая маржа ≈ 48%'),
              'ОМУ NPK 8-21: себестоимость $330/т (сырьё + производство); цена $450/т; валовая маржа ≈ 27% ($120/т)')
newh = 330 * scale
lbl = find_text(s, '$233'); set_full_text(lbl, '$330'); lbl.top = Pt(yB - newh - 16)
for sh in s.shapes:                          # green cost bar -> new height
    try:
        if sh.fill.type is not None and 'SOLID' in str(sh.fill.type) \
           and str(sh.fill.fore_color.rgb) == '5B8C6E':
            sh.top = Pt(yB - newh); sh.height = Pt(newh)
    except Exception:
        pass
set_full_text(find_text(s, '≈48%'), '≈27%')
set_full_text(find_text(s, '$217 на тонну ОМУ'), '$120 на тонну ОМУ')

# --- slide 15: divider KPI ---
s = prs.slides[14]
set_full_text(find_text(s, '≈48%'), '≈27%')
set_full_text(find_text(s, 'валовая маржа ОМУ · $233 себест. → $450 цена'),
              'валовая маржа ОМУ · $330 себест. → $450 цена')

# --- slide 19: thesis KPI + bullet ---
s = prs.slides[18]
set_full_text(find_text(s, '≈48%'), '≈27%')
set_full_text(find_text(s, 'валовая маржа ОМУ ($217/т)'), 'валовая маржа ОМУ ($120/т)')
set_full_text(find_text(s, 'Локальная себестоимость ($233/т)  защищают маржу от волатильности мировых цен.'),
              'Локальная себестоимость ($330/т) защищает маржу от волатильности мировых цен.')

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides._sldIdLst))
