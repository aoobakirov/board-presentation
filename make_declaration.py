# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# Base style
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)

def add_heading(text, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, color=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = color
    return p

def add_para(text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8, first_line=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if first_line:
        p.paragraph_format.first_line_indent = Cm(first_line)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(size)
    r.bold = bold
    return p

# Title
add_heading('ДЕКЛАРАЦИЯ', size=20, space_after=4)
add_heading('о совместных обязательствах по снижению веса', size=13, bold=False, space_after=6)
add_para('г. _____________                                                            «____» ________ 2026 г.',
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

# Preamble
add_para(
    'Мы, нижеподписавшиеся, действуя добровольно, находясь в здравом уме и твёрдой памяти, '
    'осознавая всю ответственность принимаемых на себя обязательств, а именно:',
    first_line=1.25)

add_para('•  Аубакиров Алимбек;', first_line=1.25, space_after=2)
add_para('•  Адамбаев Ильяс;', first_line=1.25, space_after=2)
add_para('•  Баймышев Азамат, —', first_line=1.25, space_after=10)

add_para(
    'далее совместно именуемые «Стороны», а по отдельности — «Сторона», заключили настоящую '
    'Декларацию (далее — «Декларация») о нижеследующем.',
    first_line=1.25, space_after=14)

# Section 1
add_heading('1. Предмет Декларации', size=13, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)
add_para(
    '1.1. Каждая из Сторон принимает на себя обязательство в срок до конца лета 2026 года '
    '(до 31 августа 2026 года включительно) снизить собственный вес на величину до 10 (десяти) '
    'килограммов и достигнуть следующих целевых показателей веса:',
    first_line=1.25)

# Table with targets
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0].cells
hdr[0].paragraphs[0].add_run('Сторона').bold = True
hdr[1].paragraphs[0].add_run('Целевой вес').bold = True

data = [
    ('Аубакиров Алимбек', '77 кг'),
    ('Адамбаев Ильяс', '83 кг'),
    ('Баймышев Азамат', '75 кг'),
]
for name, target in data:
    cells = table.add_row().cells
    cells[0].text = name
    cells[1].text = target

doc.add_paragraph().paragraph_format.space_after = Pt(6)

add_para(
    '1.2. Стороны подтверждают, что указанные целевые показатели являются для них реалистичными, '
    'безопасными для здоровья и достижимыми законными способами.',
    first_line=1.25, space_after=14)

# Section 2
add_heading('2. Ответственность Сторон', size=13, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)
add_para(
    '2.1. В случае, если какая-либо из Сторон в установленный пунктом 1.1 срок не достигнет '
    'своего целевого показателя веса, такая Сторона признаётся не исполнившей обязательство по '
    'настоящей Декларации.',
    first_line=1.25)
add_para(
    '2.2. Сторона, не исполнившая обязательство, выплачивает каждой из двух других Сторон денежную '
    'сумму в размере 500 (пятьсот) долларов США, а всего — 1 000 (одну тысячу) долларов США.',
    first_line=1.25)
add_para(
    '2.3. Стороны согласовали следующий порядок выплат при неисполнении обязательства:',
    first_line=1.25, space_after=4)
add_para(
    '•  если Аубакиров Алимбек не снизит вес — он выплачивает 500 долларов США Адамбаеву Ильясу '
    'и 500 долларов США Баймышеву Азамату;',
    first_line=1.25, space_after=2)
add_para(
    '•  если Адамбаев Ильяс не снизит вес — он выплачивает 500 долларов США Аубакирову Алимбеку '
    'и 500 долларов США Баймышеву Азамату;',
    first_line=1.25, space_after=2)
add_para(
    '•  если Баймышев Азамат не снизит вес — он выплачивает 500 долларов США Аубакирову Алимбеку '
    'и 500 долларов США Адамбаеву Ильясу.',
    first_line=1.25, space_after=8)
add_para(
    '2.4. Обязательства Сторон являются независимыми: неисполнение обязательства одной Стороной '
    'не освобождает другие Стороны от исполнения собственных обязательств. Если обязательство '
    'не исполнят несколько Сторон одновременно, каждая из них производит выплаты в порядке, '
    'установленном пунктами 2.2 и 2.3.',
    first_line=1.25)
add_para(
    '2.5. Выплата денежных средств производится в течение 14 (четырнадцати) календарных дней с '
    'момента наступления срока, указанного в пункте 1.1.',
    first_line=1.25, space_after=14)

# Section 3
add_heading('3. Заключительные положения', size=13, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)
add_para(
    '3.1. Настоящая Декларация вступает в силу с момента её подписания всеми Сторонами и действует '
    'до полного исполнения Сторонами принятых на себя обязательств.',
    first_line=1.25)
add_para(
    '3.2. Фиксация фактического веса Сторон производится путём взвешивания в присутствии всех Сторон '
    'на исправных весах. Результаты фиксируются по взаимному согласию.',
    first_line=1.25)
add_para(
    '3.3. Настоящая Декларация составлена в трёх экземплярах, имеющих одинаковую юридическую силу, '
    'по одному для каждой из Сторон.',
    first_line=1.25)
add_para(
    '3.4. Стороны заключают настоящую Декларацию по обоюдному согласию и обязуются добросовестно '
    'исполнять принятые на себя обязательства.',
    first_line=1.25, space_after=20)

# Signatures
add_heading('Подписи Сторон', size=13, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=14)

sig_table = doc.add_table(rows=0, cols=2)
sig_table.allow_autofit = True

signers = [
    'Аубакиров Алимбек',
    'Адамбаев Ильяс',
    'Баймышев Азамат',
]
for name in signers:
    row = sig_table.add_row().cells
    row[0].paragraphs[0].add_run(name)
    row[1].paragraphs[0].add_run('________________ / подпись')
    # spacer row
    spacer = sig_table.add_row().cells
    spacer[0].paragraphs[0].add_run('')

out = '/home/user/board-presentation/Декларация_о_снижении_веса.docx'
doc.save(out)
print('saved', out)
