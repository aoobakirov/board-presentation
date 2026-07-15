# -*- coding: utf-8 -*-
"""H2-2026 budget for TOO WayStar Group — formula-driven, editable drivers."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference

b=json.load(open('/home/user/board-presentation/mgmt-report/base.json'))
ENTITY='ТОО «WayStar Group» (БИН 240340014287)'
OPENCASH=113411502.27
MONTHS=['Июль','Август','Сентябрь','Октябрь','Ноябрь','Декабрь']
COLS=['C','D','E','F','G','H']  # month columns; I = total

NAVY="13294B"; BLUE="2F6690"; ORANGE="C4772F"; LGREY="EEF2F6"; GREEN="1E7A3D"; RED="B23A3A"; YELL="FFF6D5"
thin=Side(style='thin',color='C9D3DD'); border=Border(left=thin,right=thin,top=thin,bottom=thin)
NUM='#,##0;[Red]-#,##0'; PCT='0.0%'
def C(ws,r,c,v,bold=False,fill=None,fc="1A1A1A",size=10,align='left',fmt=None,italic=False,wrap=False,bd=True):
    x=ws.cell(r,c,v); x.font=Font(name='Calibri',size=size,bold=bold,color=fc,italic=italic)
    x.alignment=Alignment(horizontal=align,vertical='center',wrap_text=wrap)
    if fill: x.fill=PatternFill('solid',fgColor=fill)
    if fmt: x.number_format=fmt
    if bd: x.border=border
    return x
def banner(ws,r,ncol,text,fill=NAVY,fc="FFFFFF",size=11):
    C(ws,r,1,text,bold=True,fill=fill,fc=fc,size=size)
    for c in range(2,ncol+1): C(ws,r,c,"",fill=fill)
def hdr(ws,r,c,t): C(ws,r,c,t,bold=True,fill=NAVY,fc="FFFFFF",size=10,align='center',wrap=True)

wb=Workbook()
PRE="'Предпосылки'"

# ============================================================
# 1. ПРЕДПОСЫЛКИ (editable drivers)
# ============================================================
ws=wb.active; ws.title="Предпосылки"; ws.sheet_view.showGridLines=False
C(ws,1,1,"ПРЕДПОСЫЛКИ БЮДЖЕТА · H2 2026 (июль–декабрь)",bold=True,size=15,fc=NAVY,bd=False)
C(ws,2,1,f"{ENTITY}  ·  жёлтые ячейки — редактируемые драйверы; отчёты пересчитываются автоматически",size=9,italic=True,fc="777777",bd=False)
def drv(r,label,val,fmt=NUM,hint=""):
    C(ws,r,1,label,size=10); C(ws,r,2,val,align='right',fmt=fmt,fill=YELL,bold=True)
    C(ws,r,3,hint,size=9,italic=True,fc="777777")
banner(ws,4,3,"ВЫРУЧКА (база — факт 15.06–15.07.2026, ₸/мес)",fill=BLUE)
drv(5,"Базовая выручка — Аренда (сдача вагонов/цистерн)",round(b['rev_ar']),hint="из факт-месяца")
drv(6,"Базовая выручка — ТЭО (экспедирование, тариф)",round(b['rev_te']),hint="из факт-месяца")
drv(7,"Нормализация базы, %",0.85,PCT,"хайркат под разовые авансы QAZAQ (ТЭО 200+160 млн)")
drv(8,"Рост выручки, %/мес",0.0,PCT,"органический рост month-over-month")
C(ws,9,1,"Сезонность по месяцам (коэффициент):",size=10,italic=True)
for i,m in enumerate(MONTHS): C(ws,9,3+i,m,bold=True,size=9,align='center',fill=NAVY,fc="FFFFFF")
C(ws,10,1,"  коэффициент сезонности",size=10)
for i in range(6): C(ws,10,3+i,1.0,align='right',fmt='0.00',fill=YELL,bold=True)
banner(ws,12,3,"СЕБЕСТОИМОСТЬ (% от выручки сегмента)",fill=BLUE)
drv(13,"COGS Аренда, % от выручки Аренда",round(b['cogs_ar']/b['rev_ar'],3),PCT,"аренда вагонов у собственников")
drv(14,"COGS ТЭО, % от выручки ТЭО",round(b['cogs_te']/b['rev_te'],3),PCT,"ж/д тариф РК + иностр. транзит + ремонт")
banner(ws,16,3,"ОПЕРАЦИОННЫЕ РАСХОДЫ И НАЛОГИ",fill=BLUE)
drv(17,"OPEX базовый (ЗП, аренда, банк, связь), ₸/мес",round(b['opex']),hint="условно-постоянные")
drv(18,"Инфляция OPEX, %/мес",0.0,PCT)
drv(19,"Налоги, % от выручки (касс. оценка)",round(b['tax']/b['rev'],4),PCT,"payroll+соц.+КПН за нерезид.")
banner(ws,21,3,"ФИНАНСЫ И ДЕНЕЖНЫЕ СРЕДСТВА (₸/мес)",fill=BLUE)
drv(22,"Доход по депозитам",round(b['fin_inc']),hint="% на размещённые средства")
drv(23,"Проценты по кредитам",round(b['fin_cost']),hint="обслуживание долга (факт вкл. пени)")
drv(24,"Погашение основного долга",0,hint="плановое, если есть график")
drv(25,"Новые займы / транши",0,hint="привлечение финансирования")
drv(26,"Остаток денежных средств на начало H2",round(OPENCASH),hint="= остаток на 15.07.2026 (4 счёта)")
ws.column_dimensions['A'].width=48; ws.column_dimensions['B'].width=18
for c in 'CDEFGH': ws.column_dimensions[c].width=11
ws.column_dimensions['C'].width=12

# driver cell refs
D=dict(ar='$B$5',te='$B$6',norm='$B$7',growth='$B$8',cogs_ar='$B$13',cogs_te='$B$14',
       opex='$B$17',infl='$B$18',tax='$B$19',fininc='$B$22',fincost='$B$23',prin='$B$24',loan='$B$25',cash='$B$26')
SEAS=['$C$10','$D$10','$E$10','$F$10','$G$10','$H$10']  # seasonality per month (C10:H10)

# ============================================================
# 2. БЮДЖЕТ P&L
# ============================================================
pl=wb.create_sheet("Бюджет P&L"); pl.sheet_view.showGridLines=False
PLQ="'Бюджет P&L'"
C(pl,1,1,"БЮДЖЕТ · ОТЧЁТ О ПРИБЫЛЯХ И УБЫТКАХ · H2 2026",bold=True,size=14,fc=NAVY,bd=False)
C(pl,2,1,f"{ENTITY}  ·  ₸  ·  кассовый метод (формулы от листа «Предпосылки»)",size=10,italic=True,bd=False)
HR=4
C(pl,HR,1,"Показатель",bold=True,fill=NAVY,fc="FFFFFF"); 
for i,m in enumerate(MONTHS): hdr(pl,HR,3+i,f"{m}\n2026")
hdr(pl,HR,9,"ИТОГО H2"); hdr(pl,HR,2,"")
rows_addr={}
def prow(r,label,gen,bold=False,indent=False,color="1A1A1A",pct_of=None):
    C(pl,r,1,("      " if indent else "")+label,bold=bold,fill=LGREY if bold else None,size=10)
    C(pl,r,2,"",fill=LGREY if bold else None)
    for i,col in enumerate(COLS):
        C(pl,r,3+i,gen(i,col),bold=bold,align='right',fmt=NUM,fill=LGREY if bold else None,fc=color)
    C(pl,r,9,f"=SUM(C{r}:H{r})",bold=True,align='right',fmt=NUM,fill=LGREY,fc=color)
    rows_addr[label]=r
# month base revenue expression per segment
def rev_ar(i,col): return f"={PRE}!{D['ar']}*{PRE}!{D['norm']}*{PRE}!{SEAS[i]}*(1+{PRE}!{D['growth']})^{i}"
def rev_te(i,col): return f"={PRE}!{D['te']}*{PRE}!{D['norm']}*{PRE}!{SEAS[i]}*(1+{PRE}!{D['growth']})^{i}"
r=HR
r+=1; prow(r,"Выручка — Аренда", rev_ar, indent=True); R_ar=r
r+=1; prow(r,"Выручка — ТЭО", rev_te, indent=True); R_te=r
r+=1; prow(r,"Выручка, итого", lambda i,col:f"={col}{R_ar}+{col}{R_te}", bold=True); R_rev=r
r+=1; prow(r,"COGS — Аренда", lambda i,col:f"=-{col}{R_ar}*{PRE}!{D['cogs_ar']}", indent=True, color=RED); R_car=r
r+=1; prow(r,"COGS — ТЭО", lambda i,col:f"=-{col}{R_te}*{PRE}!{D['cogs_te']}", indent=True, color=RED); R_cte=r
r+=1; prow(r,"Себестоимость, итого", lambda i,col:f"={col}{R_car}+{col}{R_cte}", bold=True, color=RED); R_cogs=r
r+=1; prow(r,"ВАЛОВАЯ ПРИБЫЛЬ", lambda i,col:f"={col}{R_rev}+{col}{R_cogs}", bold=True, color=GREEN); R_gross=r
r+=1; prow(r,"Операционные расходы (SG&A)", lambda i,col:f"=-{PRE}!{D['opex']}*(1+{PRE}!{D['infl']})^{i}", indent=True, color=RED); R_opex=r
r+=1; prow(r,"EBITDA", lambda i,col:f"={col}{R_gross}+{col}{R_opex}", bold=True, color=GREEN); R_ebitda=r
r+=1; prow(r,"Налоги", lambda i,col:f"=-{col}{R_rev}*{PRE}!{D['tax']}", indent=True, color=RED); R_tax=r
r+=1; prow(r,"Финансовые доходы (% депозиты)", lambda i,col:f"={PRE}!{D['fininc']}", indent=True); R_fi=r
r+=1; prow(r,"Финансовые расходы (% кредиты)", lambda i,col:f"=-{PRE}!{D['fincost']}", indent=True, color=RED); R_fc=r
r+=1; prow(r,"ЧИСТАЯ ПРИБЫЛЬ", lambda i,col:f"={col}{R_ebitda}+{col}{R_tax}+{col}{R_fi}+{col}{R_fc}", bold=True, color=GREEN); R_net=r
# margin rows
r+=2; C(pl,r,1,"Справочно — маржа, %",bold=True,size=9,italic=True,bd=False)
r+=1; C(pl,r,1,"  Валовая маржа",size=9)
for i,col in enumerate(COLS): C(pl,r,3+i,f"=IFERROR({col}{R_gross}/{col}{R_rev},0)",align='right',fmt=PCT,size=9)
C(pl,r,9,f"=IFERROR(I{R_gross}/I{R_rev},0)",align='right',fmt=PCT,size=9)
r+=1; C(pl,r,1,"  EBITDA-маржа",size=9)
for i,col in enumerate(COLS): C(pl,r,3+i,f"=IFERROR({col}{R_ebitda}/{col}{R_rev},0)",align='right',fmt=PCT,size=9)
C(pl,r,9,f"=IFERROR(I{R_ebitda}/I{R_rev},0)",align='right',fmt=PCT,size=9)
r+=1; C(pl,r,1,"  Чистая маржа",size=9)
for i,col in enumerate(COLS): C(pl,r,3+i,f"=IFERROR({col}{R_net}/{col}{R_rev},0)",align='right',fmt=PCT,size=9)
C(pl,r,9,f"=IFERROR(I{R_net}/I{R_rev},0)",align='right',fmt=PCT,size=9)
pl.column_dimensions['A'].width=34; pl.column_dimensions['B'].width=2
for c in 'CDEFGH': pl.column_dimensions[c].width=14
pl.column_dimensions['I'].width=16
pl.freeze_panes="C5"

# ============================================================
# 3. БЮДЖЕТ ДДС
# ============================================================
cf=wb.create_sheet("Бюджет ДДС"); cf.sheet_view.showGridLines=False
C(cf,1,1,"БЮДЖЕТ · ДВИЖЕНИЕ ДЕНЕЖНЫХ СРЕДСТВ · H2 2026",bold=True,size=14,fc=NAVY,bd=False)
C(cf,2,1,f"{ENTITY}  ·  ₸",size=10,italic=True,bd=False)
HR=4
C(cf,HR,1,"Статья",bold=True,fill=NAVY,fc="FFFFFF"); C(cf,HR,2,"",fill=NAVY)
for i,m in enumerate(MONTHS): hdr(cf,HR,3+i,f"{m}\n2026")
hdr(cf,HR,9,"ИТОГО H2")
def cfrow(r,label,gen,bold=False,tot=True,color="1A1A1A"):
    C(cf,r,1,label,bold=bold,fill=LGREY if bold else None,size=10); C(cf,r,2,"",fill=LGREY if bold else None)
    for i,col in enumerate(COLS): C(cf,r,3+i,gen(i,col),bold=bold,align='right',fmt=NUM,fill=LGREY if bold else None,fc=color)
    if tot: C(cf,r,9,f"=SUM(C{r}:H{r})",bold=True,align='right',fmt=NUM,fill=LGREY,fc=color)
    else: C(cf,r,9,"",fill=LGREY if bold else None)
r=HR
r+=1; R_open=r
C(cf,r,1,"Денежные средства на начало месяца",bold=True,size=10)
C(cf,r,3,f"={PRE}!{D['cash']}",align='right',fmt=NUM,bold=True)
for i in range(1,6): C(cf,r,3+i,f"={COLS[i-1]}{r+6}",align='right',fmt=NUM,bold=True)  # prev ending
C(cf,r,9,f"=C{r}",align='right',fmt=NUM,bold=True)
r+=1; cfrow(r,"Чистая прибыль (операц. поток)", lambda i,col:f"={PLQ}!{col}{R_net}"); R_ocf=r
r+=1; cfrow(r,"Погашение основного долга", lambda i,col:f"=-{PRE}!{D['prin']}", color=RED)
Rp=r
r+=1; cfrow(r,"Новые займы / транши", lambda i,col:f"={PRE}!{D['loan']}"); Rl=r
r+=1; cfrow(r,"Чистый денежный поток за месяц", lambda i,col:f"={col}{R_ocf}+{col}{Rp}+{col}{Rl}", bold=True, color=GREEN); R_ncf=r
r+=1; R_end=r
C(cf,r,1,"Денежные средства на конец месяца",bold=True,size=10,fill=LGREY)
C(cf,r,2,"",fill=LGREY)
for i,col in enumerate(COLS): C(cf,r,3+i,f"={col}{R_open}+{col}{R_ncf}",align='right',fmt=NUM,bold=True,fill=LGREY,fc=GREEN)
C(cf,r,9,f"=H{r}",align='right',fmt=NUM,bold=True,fill=LGREY,fc=GREEN)
# ensure R_end == r+... : R_open is r-6? we referenced r+6 for prev ending; verify
r+=2
C(cf,r,1,"Свободный денежный поток (для депозитов/дивидендов)",bold=True,size=10,italic=True)
for i,col in enumerate(COLS): C(cf,r,3+i,f"={col}{R_ocf}+{col}{Rp}",align='right',fmt=NUM,italic=True)
C(cf,r,9,f"=SUM(C{r}:H{r})",align='right',fmt=NUM,bold=True)
cf.column_dimensions['A'].width=44; cf.column_dimensions['B'].width=2
for c in 'CDEFGH': cf.column_dimensions[c].width=14
cf.column_dimensions['I'].width=16
cf.freeze_panes="C5"

# fix opening-month links: month i opening = previous month ending (R_end)
for i in range(1,6):
    C(cf,R_open,3+i,f"={COLS[i-1]}{R_end}",align='right',fmt=NUM,bold=True)

# ============================================================
# 4. СЦЕНАРИИ (computed static)
# ============================================================
def h2(norm,growth,car,cte,tax=b['tax']/b['rev'],opex_infl=0.0):
    rev=eb=net=0; cash=OPENCASH
    for i in range(6):
        f=norm*(1+growth)**i
        ra=b['rev_ar']*f; rt=b['rev_te']*f; rr=ra+rt
        cogs=ra*car+rt*cte; g=rr-cogs
        op=b['opex']*(1+opex_infl)**i; e=g-op
        tx=rr*tax; nt=e-tx+b['fin_inc']-b['fin_cost']
        rev+=rr; eb+=e; net+=nt; cash+=nt
    return rev,eb,net,cash
car=b['cogs_ar']/b['rev_ar']; cte=b['cogs_te']/b['rev_te']
scen=[
 ("Консервативный", "норм. 70%, COGS +2 п.п., рост 0%", h2(0.70,0.0,car+0.02,cte+0.02)),
 ("Базовый (рекоменд.)", "норм. 85%, COGS факт, рост 0%", h2(0.85,0.0,car,cte)),
 ("Оптимистичный", "норм. 100%, рост +2%/мес, COGS −1 п.п.", h2(1.00,0.02,car-0.01,cte-0.01)),
]
sc=wb.create_sheet("Сценарии"); sc.sheet_view.showGridLines=False
C(sc,1,1,"СЦЕНАРИИ БЮДЖЕТА · H2 2026",bold=True,size=14,fc=NAVY,bd=False)
C(sc,2,1,"Чувствительность к нормализации выручки (риск концентрации QAZAQ-ASTYQ) и марже",size=10,italic=True,bd=False)
HR=4
for c,t in [(1,"Сценарий"),(2,"Предпосылки"),(3,"Выручка H2"),(4,"EBITDA H2"),(5,"Чистая приб. H2"),(6,"Остаток ден. на конец")]: hdr(sc,HR,c,t)
for i,(nm,desc,(rev,eb,net,cash)) in enumerate(scen):
    r=HR+1+i
    col=GREEN if 'Базов' in nm else ("1A1A1A")
    C(sc,r,1,nm,bold=True,size=10,fc=col); C(sc,r,2,desc,size=9,italic=True)
    C(sc,r,3,round(rev),align='right',fmt=NUM); C(sc,r,4,round(eb),align='right',fmt=NUM)
    C(sc,r,5,round(net),align='right',fmt=NUM); C(sc,r,6,round(cash),align='right',fmt=NUM)
r=HR+1+len(scen)+1
C(sc,r,1,"Метрики (базовый):",bold=True,size=10,bd=False)
bexp=scen[1][2]
for lbl,val,fmt in [("Выручка H2",bexp[0],NUM),("EBITDA-маржа",bexp[1]/bexp[0],PCT),("Чистая маржа",bexp[2]/bexp[0],PCT),
                    ("Средняя чистая прибыль/мес",bexp[2]/6,NUM),("Прирост денежных средств за H2",bexp[3]-OPENCASH,NUM)]:
    r+=1; C(sc,r,1,"   "+lbl,size=10); C(sc,r,3,round(val) if fmt==NUM else val,align='right',fmt=fmt)
sc.column_dimensions['A'].width=22; sc.column_dimensions['B'].width=40
for c in 'CDEF': sc.column_dimensions[c].width=18

# ============================================================
# 5. БАЗА (факт) — traceability
# ============================================================
bs=wb.create_sheet("База (факт)"); bs.sheet_view.showGridLines=False
C(bs,1,1,"БАЗА БЮДЖЕТА — ФАКТ ЗА 15.06–15.07.2026 (₸/мес)",bold=True,size=13,fc=NAVY,bd=False)
C(bs,2,1,"Месячный run-rate из консолидированной управленческой отчётности (4 счёта)",size=10,italic=True,bd=False)
rowsb=[("Выручка — Аренда",b['rev_ar']),("Выручка — ТЭО",b['rev_te']),("Выручка, итого",b['rev'],True),
 ("COGS — Аренда",-b['cogs_ar']),("COGS — ТЭО",-b['cogs_te']),("Себестоимость, итого",-b['cogs'],True),
 ("ВАЛОВАЯ ПРИБЫЛЬ",b['gross'],True),("Операционные расходы",-b['opex']),("EBITDA",b['ebitda'],True),
 ("Налоги",-b['tax']),("Финансовые доходы",b['fin_inc']),("Финансовые расходы",-b['fin_cost']),
 ("ЧИСТАЯ ПРИБЫЛЬ",b['net'],True)]
r=4; hdr(bs,r,1,"Показатель"); hdr(bs,r,2,"₸/мес"); hdr(bs,r,3,"% выручки")
for it in rowsb:
    r+=1; lbl=it[0]; val=it[1]; bold=len(it)>2
    C(bs,r,1,lbl,bold=bold,fill=LGREY if bold else None,size=10)
    C(bs,r,2,round(val),bold=bold,align='right',fmt=NUM,fill=LGREY if bold else None,fc=(GREEN if bold and val>0 else ("1A1A1A")))
    C(bs,r,3,val/b['rev'],bold=bold,align='right',fmt=PCT,fill=LGREY if bold else None)
r+=2
C(bs,r,1,"Оговорка: база — 1 месяц. Крупные авансовые платежи QAZAQ-ASTYQ (ТЭО «предоставление цистерн» 200 и 160 млн ₸) могут покрывать несколько месяцев, поэтому в бюджете применяется нормализация (по умолчанию 85%). Отчётность кассовая. Рекомендуется калибровать базовую выручку по договорам/нескольким месяцам факта.",size=9,italic=True,fc="777777",wrap=True,bd=False)
bs.merge_cells(start_row=r,start_column=1,end_row=r+3,end_column=3)
bs.column_dimensions['A'].width=40; bs.column_dimensions['B'].width=18; bs.column_dimensions['C'].width=12

# ============================================================
# CHARTS on a dashboard-ish area of Сценарии
# ============================================================
# monthly revenue & EBITDA from P&L (line), cash from ДДС (line) — reference formula cells
rc=LineChart(); rc.title="Выручка и EBITDA по месяцам, ₸"; rc.height=7; rc.width=14; rc.style=2
data=Reference(pl,min_col=3,max_col=8,min_row=R_rev,max_row=R_rev)
data2=Reference(pl,min_col=3,max_col=8,min_row=R_ebitda,max_row=R_ebitda)
cats=Reference(pl,min_col=3,max_col=8,min_row=HR,max_row=HR)
rc.add_data(data,from_rows=True,titles_from_data=False); rc.add_data(data2,from_rows=True,titles_from_data=False)
rc.set_categories(cats); rc.series[0].graphicalProperties.line.solidFill=NAVY; rc.series[1].graphicalProperties.line.solidFill=ORANGE
sc.add_chart(rc,"A16")
cc=LineChart(); cc.title="Прогноз остатка денежных средств, ₸"; cc.height=7; cc.width=14; cc.style=2
cdata=Reference(cf,min_col=3,max_col=8,min_row=R_end,max_row=R_end)
cc.add_data(cdata,from_rows=True,titles_from_data=False); cc.set_categories(cats)
cc.series[0].graphicalProperties.line.solidFill=GREEN
sc.add_chart(cc,"A32")

# order
order=["Сценарии","Бюджет P&L","Бюджет ДДС","Предпосылки","База (факт)"]
wb._sheets.sort(key=lambda s: order.index(s.title))
wb.active=0
out='/home/user/board-presentation/mgmt-report/WayStar_Бюджет_H2_2026.xlsx'
wb.save(out)
print("SAVED:",out)
for nm,desc,(rev,eb,net,cash) in scen:
    print(f"  {nm:22} выр {rev/1e9:.2f} млрд | EBITDA {eb/1e6:,.0f} млн | чист {net/1e6:,.0f} млн | ост {cash/1e6:,.0f} млн")
