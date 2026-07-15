# -*- coding: utf-8 -*-
"""Build management report Excel — FORMULA-DRIVEN (SUMIFS over register)."""
import json, re, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

rows=json.load(open('/home/user/board-presentation/mgmt-report/ledger_classified.json'))
BAL={'БЦК':(75346.24,110840346.22),'Halyk':(222587.02,2526900.00),'RBK':(0.0,44256.05)}
PERIOD="15.06.2026 – 15.07.2026"; ENTITY='ТОО «WayStar Group» (БИН 240340014287)'

def norm(n):
    n=re.sub(r'^(Товарищество с ограниченной ответственностью|Акционерное общество|ТОО|АО|Индивидуальный предприниматель|ИП|Частная компания|Производственный кооператив)\s*','',n,flags=re.I)
    return n.strip(' "«»') or '(без наименования)'
for x in rows: x['cpn']=norm(x['cp'])

# client & category lists (order only; values via formulas)
rev_by=collections.defaultdict(float)
for x in rows:
    if x['flow']=='REVENUE': rev_by[x['cpn']]+=x['credit']
clients=[k for k,_ in sorted(rev_by.items(),key=lambda kv:-kv[1])]
cost_by=collections.defaultdict(float)
for x in rows:
    if x['flow'] in ('COGS','OPEX'): cost_by[x['cpn']]+=x['debit']
suppliers=[k for k,_ in sorted(cost_by.items(),key=lambda kv:-kv[1])][:20]
cogs_cats=list(dict.fromkeys(x['cat'] for x in rows if x['flow']=='COGS'))
opex_cats=[c for c,_ in sorted(((c,sum(x['debit'] for x in rows if x['flow']=='OPEX' and x['cat']==c))
            for c in dict.fromkeys(x['cat'] for x in rows if x['flow']=='OPEX')),key=lambda kv:-kv[1])]

# ---- styles ----
NAVY="13294B"; BLUE="2F6690"; LGREY="EEF2F6"; GREEN="1E7A3D"; RED="B23A3A"
thin=Side(style='thin', color='C9D3DD'); border=Border(left=thin,right=thin,top=thin,bottom=thin)
NUM='#,##0;[Red]-#,##0'; PCT='0.0%'
def cell(ws,r,c,v,bold=False,fill=None,fontcolor="1A1A1A",size=10,align='left',fmt=None,italic=False,wrap=False,bd=True):
    x=ws.cell(r,c,v); x.font=Font(name='Calibri',size=size,bold=bold,color=fontcolor,italic=italic)
    x.alignment=Alignment(horizontal=align,vertical='center',wrap_text=wrap)
    if fill: x.fill=PatternFill('solid',fgColor=fill)
    if fmt: x.number_format=fmt
    if bd: x.border=border
    return x
def banner(ws,r,ncol,text,fill=NAVY,fc="FFFFFF",size=11):
    cell(ws,r,1,text,bold=True,fill=fill,fontcolor=fc,size=size)
    for c in range(2,ncol+1): cell(ws,r,c,"",fill=fill)
def hdr(ws,r,c,t): cell(ws,r,c,t,bold=True,fill=NAVY,fontcolor="FFFFFF",size=10,align='center',wrap=True)
def note(ws,r,text,ncol,span=2):
    cell(ws,r,1,text,italic=True,size=9,fontcolor="777777",wrap=True,bd=False)
    ws.merge_cells(start_row=r,start_column=1,end_row=r+span,end_column=ncol)

wb=Workbook()

# ============================================================
# REGISTER FIRST (data source for all SUMIFS)
# ============================================================
reg=wb.active; reg.title="Реестр операций"; reg.sheet_view.showGridLines=False
cell(reg,1,1,"РЕЕСТР ОПЕРАЦИЙ (классифицированный) — источник данных для формул",bold=True,size=12,fontcolor=NAVY,bd=False)
HROW=3
for c,t in enumerate(["Банк","Дата","Контрагент","Контрагент (норм.)","БИН","Поступление","Выплата","КНП","Тип","Категория","Направление","Назначение платежа"],1):
    hdr(reg,HROW,c,t)
data=sorted(rows,key=lambda x:(x['date'][6:10]+x['date'][3:5]+x['date'][0:2], x['bank']))
r=HROW
for x in data:
    r+=1
    cell(reg,r,1,x['bank'],size=9); cell(reg,r,2,x['date'],size=9,align='center')
    cell(reg,r,3,x['cp'][:45],size=9); cell(reg,r,4,x['cpn'],size=9); cell(reg,r,5,x.get('bin',''),size=9)
    cell(reg,r,6,x['credit'] or None,size=9,align='right',fmt=NUM)
    cell(reg,r,7,x['debit'] or None,size=9,align='right',fmt=NUM)
    cell(reg,r,8,x['knp'],size=9,align='center'); cell(reg,r,9,x['flow'],size=9)
    cell(reg,r,10,x['cat'],size=9); cell(reg,r,11,x['line'],size=9,align='center')
    cell(reg,r,12,x['purpose'][:90],size=9)
LAST=r
for i,w in enumerate([8,11,32,30,14,14,14,7,11,30,12,55],1): reg.column_dimensions[get_column_letter(i)].width=w
reg.freeze_panes="A4"; reg.auto_filter.ref=f"A{HROW}:L{LAST}"

# SUMIFS reference ranges (full columns)
R="'Реестр операций'"
CR=f"{R}!$F:$F"; DB=f"{R}!$G:$G"; NM=f"{R}!$D:$D"; FL=f"{R}!$I:$I"; CT=f"{R}!$J:$J"; LN=f"{R}!$K:$K"
def sifs(val_rng, *pairs):
    s=val_rng
    for rng,crit in pairs: s+=f",{rng},{crit}"
    return f"=SUMIFS({s})"
def q(x): return '"'+x.replace('"','""')+'"'

# ============================================================
# 1. P&L (кассовый) — build first, others reference it
# ============================================================
pl=wb.create_sheet("P&L (кассовый)"); pl.sheet_view.showGridLines=False
PLQ="'P&L (кассовый)'"
cell(pl,1,1,"ОТЧЁТ О ПРИБЫЛЯХ И УБЫТКАХ (кассовый метод)",bold=True,size=14,fontcolor=NAVY,bd=False)
cell(pl,2,1,f"{ENTITY}  ·  {PERIOD}  ·  ₸",size=10,italic=True,bd=False)
r=4
for c,t in [(1,"Показатель"),(2,"Сумма, ₸"),(3,"% к чист. выручке")]: hdr(pl,r,c,t)
addr={}
def plrow(label,formula,key=None,bold=False,indent=False,color="1A1A1A"):
    global r; r+=1
    cell(pl,r,1,("      " if indent else "")+label,bold=bold,fill=LGREY if bold else None,size=10)
    cell(pl,r,2,formula,bold=bold,align='right',fmt=NUM,fill=LGREY if bold else None,fontcolor=color)
    if key: addr[key]=r
    return r
# revenue
plrow("Выручка от реализации услуг", sifs(CR,(FL,q("REVENUE"))), key='rev', bold=True)
plrow("минус: возврат средств клиентам", "=-"+sifs(DB,(FL,q("CONTRA")))[1:], key='refund', indent=True)
plrow("Чистая выручка", f"=B{addr['rev']}+B{addr['refund']}", key='netrev', bold=True)
# cogs
for cat in cogs_cats:
    plrow("минус: "+cat, "=-"+sifs(DB,(FL,q("COGS")),(CT,q(cat)))[1:], indent=True)
plrow("Себестоимость услуг, итого", "=-"+sifs(DB,(FL,q("COGS")))[1:], key='cogs', bold=True, color=RED)
plrow("ВАЛОВАЯ ПРИБЫЛЬ", f"=B{addr['netrev']}+B{addr['cogs']}", key='gross', bold=True, color=GREEN)
# opex
for cat in opex_cats:
    plrow("минус: "+cat, "=-"+sifs(DB,(FL,q("OPEX")),(CT,q(cat)))[1:], indent=True)
plrow("Операционные расходы, итого", "=-"+sifs(DB,(FL,q("OPEX")))[1:], key='opex', bold=True, color=RED)
plrow("EBITDA", f"=B{addr['gross']}+B{addr['opex']}", key='ebitda', bold=True, color=GREEN)
plrow("минус: налоги и социальные платежи", "=-"+sifs(DB,(FL,q("TAX")))[1:], key='tax', indent=True)
plrow("плюс: финансовые доходы (% по депозитам)", f"={sifs(CR,(FL,q('FIN_INCOME')))[1:]}+{sifs(CR,(FL,q('FIN_COST')))[1:]}", key='finin', indent=True)
plrow("минус: финансовые расходы (% и пени)", "=-"+sifs(DB,(FL,q("FIN_COST")))[1:], key='fincost', indent=True)
plrow("ЧИСТАЯ ПРИБЫЛЬ", f"=B{addr['ebitda']}+B{addr['tax']}+B{addr['finin']}+B{addr['fincost']}", key='net', bold=True, color=GREEN)
# pct column
for k,rr in addr.items():
    isbold = k in ('rev','netrev','cogs','gross','opex','ebitda','net')
    cell(pl,rr,3,f"=IFERROR(B{rr}/$B${addr['netrev']},0)",bold=isbold,align='right',fmt=PCT,fill=LGREY if isbold else None)
r+=2
note(pl,r,"Кассовый метод по банковским выпискам: «выручка» = поступления от клиентов, «расходы» = фактические платежи за период. Не учитывает начисления, дебиторскую/кредиторскую задолженность, амортизацию, НДС к зачёту/уплате. Депозиты, займы и переводы между своими счетами исключены из P&L (см. ДДС). Все суммы — формулы SUMIFS по листу «Реестр операций».",3,2)
pl.column_dimensions['A'].width=54; pl.column_dimensions['B'].width=18; pl.column_dimensions['C'].width=15

# ============================================================
# 2. P&L ПО КЛИЕНТАМ (Аренда / ТЭО), formula-driven
# ============================================================
ws=wb.create_sheet("P&L по клиентам"); ws.sheet_view.showGridLines=False
cell(ws,1,1,"P&L В РАЗБИВКЕ ПО КЛИЕНТАМ — Аренда и ТЭО",bold=True,size=14,fontcolor=NAVY,bd=False)
cell(ws,2,1,f"{ENTITY}  ·  {PERIOD}  ·  ₸",size=10,italic=True,bd=False)
note(ws,3,"Выручка по клиентам — фактическая, из назначений платежей разнесена на «Аренду» (сдача вагонов/цистерн) и «ТЭО» (экспедирование, ж/д тариф, предоставление ПС). Прямые расходы распределены по клиентам пропорционально их выручке в соответствующем направлении; накладные — пропорционально общей выручке (в выписках нет привязки платежей поставщикам к клиенту).",11,1)
HR=5
heads=["Клиент","Выручка: Аренда","Выручка: ТЭО","Выручка ИТОГО","Расходы: Аренда","Расходы: ТЭО","Прямые расходы ИТОГО","Валовая прибыль","GM %","Накладные (распр.)","Чистая прибыль"]
for c,t in enumerate(heads,1): hdr(ws,HR,c,t)
TOT=HR+len(clients)+1   # ИТОГО row number (known in advance)
# client rows
for i,name in enumerate(clients):
    r=HR+1+i
    cell(ws,r,1,name,size=10)
    cell(ws,r,2, sifs(CR,(NM,f"$A{r}"),(FL,q("REVENUE")),(LN,q("Аренда"))), align='right',fmt=NUM)
    cell(ws,r,3, sifs(CR,(NM,f"$A{r}"),(FL,q("REVENUE")),(LN,q("ТЭО"))), align='right',fmt=NUM)
    cell(ws,r,4, f"=B{r}+C{r}", align='right',fmt=NUM)
    cell(ws,r,5, f"=IFERROR($E${TOT}*B{r}/$B${TOT},0)", align='right',fmt=NUM,fontcolor=RED)
    cell(ws,r,6, f"=IFERROR($F${TOT}*C{r}/$C${TOT},0)", align='right',fmt=NUM,fontcolor=RED)
    cell(ws,r,7, f"=E{r}+F{r}", align='right',fmt=NUM,fontcolor=RED)
    cell(ws,r,8, f"=D{r}-G{r}", align='right',fmt=NUM)
    cell(ws,r,9, f"=IFERROR(H{r}/D{r},0)", align='right',fmt=PCT)
    cell(ws,r,10,f"=IFERROR($J${TOT}*D{r}/$D${TOT},0)", align='right',fmt=NUM,fontcolor=RED)
    cell(ws,r,11,f"=H{r}-J{r}", align='right',fmt=NUM,fontcolor=GREEN)
# TOTAL row (actual SUMIFS totals)
r=TOT
cell(ws,r,1,"ИТОГО",bold=True,fill=LGREY)
cell(ws,r,2, sifs(CR,(FL,q("REVENUE")),(LN,q("Аренда"))), bold=True,align='right',fmt=NUM,fill=LGREY)
cell(ws,r,3, sifs(CR,(FL,q("REVENUE")),(LN,q("ТЭО"))), bold=True,align='right',fmt=NUM,fill=LGREY)
cell(ws,r,4, f"=B{r}+C{r}", bold=True,align='right',fmt=NUM,fill=LGREY)
cell(ws,r,5, sifs(DB,(FL,q("COGS")),(LN,q("Аренда"))), bold=True,align='right',fmt=NUM,fill=LGREY,fontcolor=RED)
cell(ws,r,6, sifs(DB,(FL,q("COGS")),(LN,q("ТЭО"))), bold=True,align='right',fmt=NUM,fill=LGREY,fontcolor=RED)
cell(ws,r,7, f"=E{r}+F{r}", bold=True,align='right',fmt=NUM,fill=LGREY,fontcolor=RED)
cell(ws,r,8, f"=D{r}-G{r}", bold=True,align='right',fmt=NUM,fill=LGREY,fontcolor=GREEN)
cell(ws,r,9, f"=IFERROR(H{r}/D{r},0)", bold=True,align='right',fmt=PCT,fill=LGREY)
# overhead total = OPEX+TAX+FIN_COST(debit)+CONTRA - FIN_INCOME - FIN_COST(credit)
oh=f"{sifs(DB,(FL,q('OPEX')))[1:]}+{sifs(DB,(FL,q('TAX')))[1:]}+{sifs(DB,(FL,q('FIN_COST')))[1:]}+{sifs(DB,(FL,q('CONTRA')))[1:]}-{sifs(CR,(FL,q('FIN_INCOME')))[1:]}-{sifs(CR,(FL,q('FIN_COST')))[1:]}"
cell(ws,r,10,"="+oh, bold=True,align='right',fmt=NUM,fill=LGREY,fontcolor=RED)
cell(ws,r,11,f"=H{r}-J{r}", bold=True,align='right',fmt=NUM,fill=LGREY,fontcolor=GREEN)
ws.column_dimensions['A'].width=40
for col in 'BCDEFGHK': ws.column_dimensions[col].width=15
ws.column_dimensions['I'].width=8; ws.column_dimensions['J'].width=15
ws.freeze_panes="B6"

# ============================================================
# 3. РЕЗЮМЕ (references P&L cells)
# ============================================================
ws=wb.create_sheet("Резюме"); ws.sheet_view.showGridLines=False
cell(ws,1,1,"УПРАВЛЕНЧЕСКАЯ ОТЧЁТНОСТЬ",bold=True,size=16,fontcolor=NAVY,bd=False)
cell(ws,2,1,ENTITY,size=11,fontcolor=BLUE,bd=False)
cell(ws,3,1,f"Период: {PERIOD}  ·  выписки 3 банков (БЦК, Halyk, Bank RBK)  ·  валюта: KZT",size=10,italic=True,bd=False)
cell(ws,4,1,"Метод: кассовый (по фактическому движению денег). Все показатели — формулы.",size=9,italic=True,fontcolor="777777",bd=False)
r=6; banner(ws,r,4,"КЛЮЧЕВЫЕ ПОКАЗАТЕЛИ (₸)")
def kpi(label,formula,big=False,col="1A1A1A",pct=False):
    global r; r+=1
    cell(ws,r,1,label,bold=big,size=11 if big else 10,fill=LGREY if big else None)
    for c in (2,3): cell(ws,r,c,"",fill=LGREY if big else None)
    cell(ws,r,4,formula,bold=big,align='right',fmt=(PCT if pct else NUM),fontcolor=col,size=11 if big else 10,fill=LGREY if big else None)
P=lambda k: f"{PLQ}!B{addr[k]}"
kpi("Выручка (поступления от клиентов)", f"={P('rev')}")
kpi("Себестоимость (ж/д тариф, вагоны, ремонт)", f"={P('cogs')}", col=RED)
kpi("ВАЛОВАЯ ПРИБЫЛЬ", f"={P('gross')}", big=True, col=GREEN)
kpi("   валовая маржа, %", f"=IFERROR({P('gross')}/{P('netrev')},0)", pct=True)
kpi("Операционные расходы (SG&A)", f"={P('opex')}", col=RED)
kpi("EBITDA", f"={P('ebitda')}", big=True, col=GREEN)
kpi("Налоги уплаченные", f"={P('tax')}", col=RED)
kpi("Финансовые доходы / расходы", f"={P('finin')}+{P('fincost')}")
kpi("ЧИСТАЯ ПРИБЫЛЬ (кассовая)", f"={P('net')}", big=True, col=GREEN)
kpi("   чистая маржа, %", f"=IFERROR({P('net')}/{P('netrev')},0)", pct=True)
r+=1; banner(ws,r,4,"ДВИЖЕНИЕ ДЕНЕЖНЫХ СРЕДСТВ (свод, ₸)")
open_bal=sum(v[0] for v in BAL.values()); close_bal=sum(v[1] for v in BAL.values())
DDS="'ДДС (Cash Flow)'"
def cfl(label,formula,tot=False):
    global r; r+=1
    cell(ws,r,1,label,bold=tot,fill=LGREY if tot else None,size=10)
    for c in (2,3): cell(ws,r,c,"",fill=LGREY if tot else None)
    cell(ws,r,4,formula,bold=tot,align='right',fmt=NUM,fill=LGREY if tot else None,fontcolor=(GREEN if tot else "1A1A1A"))
cfl("Денежные средства на начало периода", f"={open_bal}", tot=True)
cfl("Операционный денежный поток", f"={P('net')}")
cfl("Инвестиц./казначейские (депозиты, валюта)", f"={sifs(CR,(FL,q('TREASURY')))[1:]}-{sifs(DB,(FL,q('TREASURY')))[1:]}")
cfl("Финансовые потоки (займы, кредиты)", f"={sifs(CR,(FL,q('FINANCING')))[1:]}-{sifs(DB,(FL,q('FINANCING')))[1:]}")
cfl("Переводы на прочие счета компании (КЛС и пр.)", f"={sifs(CR,(FL,q('INTERNAL')))[1:]}-{sifs(DB,(FL,q('INTERNAL')))[1:]}")
cfl("Прочее", f"={sifs(CR,(FL,q('OTHER_IN')))[1:]}")
cfl("Денежные средства на конец периода", f"={close_bal}", tot=True)
ws.column_dimensions['A'].width=48
for c in 'BCD': ws.column_dimensions[c].width=16

# ============================================================
# 4. ДДС (Cash Flow) — SUMIFS
# ============================================================
ws=wb.create_sheet("ДДС (Cash Flow)"); ws.sheet_view.showGridLines=False
cell(ws,1,1,"ОТЧЁТ О ДВИЖЕНИИ ДЕНЕЖНЫХ СРЕДСТВ (прямой метод)",bold=True,size=14,fontcolor=NAVY,bd=False)
cell(ws,2,1,f"{ENTITY}  ·  {PERIOD}  ·  ₸",size=10,italic=True,bd=False)
r=4
for c,t in [(1,"Статья"),(2,"Поступления"),(3,"Выплаты"),(4,"Итого")]: hdr(ws,r,c,t)
def cf(label,inf,outf,bold=False):
    global r; r+=1
    cell(ws,r,1,("" if bold else "      ")+label,bold=bold,fill=LGREY if bold else None,size=10)
    cell(ws,r,2,inf,align='right',fmt=NUM,fill=LGREY if bold else None)
    cell(ws,r,3,outf,align='right',fmt=NUM,fill=LGREY if bold else None,fontcolor=RED)
    cell(ws,r,4,f"=N({chr(78)})" if False else f"=IFERROR(B{r},0)-IFERROR(-C{r},0)",align='right',fmt=NUM,bold=bold,fill=LGREY if bold else None)
    return r
def cf2(label,inf,outf,bold=False):
    global r; r+=1
    cell(ws,r,1,("" if bold else "      ")+label,bold=bold,fill=LGREY if bold else None,size=10)
    cell(ws,r,2,inf if inf else None,align='right',fmt=NUM,fill=LGREY if bold else None)
    cell(ws,r,3,("=-"+outf[1:]) if outf else None,align='right',fmt=NUM,fill=LGREY if bold else None,fontcolor=RED)
    cell(ws,r,4,f"=IFERROR(B{r},0)+IFERROR(C{r},0)",align='right',fmt=NUM,bold=bold,fill=LGREY if bold else None)
    return r
cf2("Остаток денежных средств на начало периода", f"={sum(v[0] for v in BAL.values())}", None, bold=True)
r+=1; banner(ws,r,4,"ОПЕРАЦИОННАЯ ДЕЯТЕЛЬНОСТЬ",fill=BLUE)
o1=cf2("Поступления от клиентов", sifs(CR,(FL,q("REVENUE"))), None)
o2=cf2("Возврат средств клиентам", None, sifs(DB,(FL,q("CONTRA"))))
o3=cf2("Оплата перевозчикам / ж/д тариф / вагоны (COGS)", None, sifs(DB,(FL,q("COGS"))))
o4=cf2("Операционные расходы (ЗП, аренда, банк, связь)", None, sifs(DB,(FL,q("OPEX"))))
o5=cf2("Налоги и социальные платежи", None, sifs(DB,(FL,q("TAX"))))
o6=cf2("Проценты и пени по кредитам", None, sifs(DB,(FL,q("FIN_COST"))))
o7=cf2("Проценты по депозитам", sifs(CR,(FL,q("FIN_INCOME"))), None)
osum=cf2("Итого по операционной деятельности", None, None, bold=True)
cell(ws,osum,4,f"=SUM(D{o1}:D{o7})",bold=True,align='right',fmt=NUM,fill=LGREY,fontcolor=GREEN)
r+=1; banner(ws,r,4,"ИНВЕСТИЦИОННАЯ / КАЗНАЧЕЙСКАЯ ДЕЯТЕЛЬНОСТЬ",fill=BLUE)
i1=cf2("Снятие / возврат депозитов", sifs(CR,(CT,q("Снятие/возврат депозита"))), None)
i2=cf2("Размещение депозитов", None, sifs(DB,(CT,q("Размещение депозита"))))
i3=cf2("Покупка иностранной валюты (FX)", None, sifs(DB,(CT,q("Покупка валюты (FX)"))))
isum=cf2("Итого по инвестиц./казнач. деятельности", None, None, bold=True)
cell(ws,isum,4,f"=SUM(D{i1}:D{i3})",bold=True,align='right',fmt=NUM,fill=LGREY)
r+=1; banner(ws,r,4,"ФИНАНСОВАЯ ДЕЯТЕЛЬНОСТЬ",fill=BLUE)
f1=cf2("Займы полученные / возврат выданных займов", sifs(CR,(FL,q("FINANCING"))), None)
f2=cf2("Погашение и выдача займов", None, sifs(DB,(FL,q("FINANCING"))))
fsum=cf2("Итого по финансовой деятельности", None, None, bold=True)
cell(ws,fsum,4,f"=SUM(D{f1}:D{f2})",bold=True,align='right',fmt=NUM,fill=LGREY)
r+=1; banner(ws,r,4,"ВНУТРЕННИЕ ПЕРЕВОДЫ И ПРОЧЕЕ",fill=BLUE)
n1=cf2("Переводы между собственными счетами (в т.ч. КЛС)", sifs(CR,(FL,q("INTERNAL"))), sifs(DB,(FL,q("INTERNAL"))))
n2=cf2("Прочие поступления", sifs(CR,(FL,q("OTHER_IN"))), None)
r+=1; banner(ws,r,4,"ЧИСТОЕ ИЗМЕНЕНИЕ ДЕНЕЖНЫХ СРЕДСТВ")
chg=r
cell(ws,r,4,f"=D{osum}+D{isum}+D{fsum}+D{n1}+D{n2}",bold=True,align='right',fmt=NUM,fill=NAVY,fontcolor="FFFFFF")
endrow=cf2("Остаток денежных средств на конец периода", None, None, bold=True)
cell(ws,endrow,4,f"=D5+D{chg}",bold=True,align='right',fmt=NUM,fill=LGREY,fontcolor=GREEN)  # D5 = opening
cell(ws,endrow,2,f"={sum(v[1] for v in BAL.values())}",bold=True,align='right',fmt=NUM,fill=LGREY)
r=endrow+2
cell(ws,r,1,"Сверка по банковским счетам (контроль):",bold=True,size=10,bd=False); r+=1
for c,t in [(1,"Банк / счёт"),(2,"Остаток на начало"),(3,"Остаток на конец")]: hdr(ws,r,c,t)
cell(ws,r,4,"",fill=NAVY)
names={'БЦК':'АО «Банк ЦентрКредит»','Halyk':'АО «Народный Банк»','RBK':'АО «Bank RBK»'}
b0=r
for b,(o,cl) in BAL.items():
    r+=1; cell(ws,r,1,names[b],size=10); cell(ws,r,2,o,align='right',fmt=NUM); cell(ws,r,3,cl,align='right',fmt=NUM); cell(ws,r,4,"")
r+=1
cell(ws,r,1,"ИТОГО по 3 счетам",bold=True,fill=LGREY)
cell(ws,r,2,f"=SUM(B{b0+1}:B{r-1})",bold=True,align='right',fmt=NUM,fill=LGREY)
cell(ws,r,3,f"=SUM(C{b0+1}:C{r-1})",bold=True,align='right',fmt=NUM,fill=LGREY); cell(ws,r,4,"",fill=LGREY)
ws.column_dimensions['A'].width=54
for c in 'BCD': ws.column_dimensions[c].width=18

# ============================================================
# 5. ВЫРУЧКА ПО КЛИЕНТАМ
# ============================================================
ws=wb.create_sheet("Выручка по клиентам"); ws.sheet_view.showGridLines=False
cell(ws,1,1,"ВЫРУЧКА ПО КЛИЕНТАМ",bold=True,size=14,fontcolor=NAVY,bd=False)
cell(ws,2,1,f"{PERIOD}  ·  ₸",size=10,italic=True,bd=False)
HR=4
for c,t in [(1,"Клиент"),(2,"Аренда"),(3,"ТЭО"),(4,"Операций"),(5,"Сумма ИТОГО"),(6,"Доля")]: hdr(ws,HR,c,t)
TOTv=HR+len(clients)+1
for i,name in enumerate(clients):
    r=HR+1+i
    cell(ws,r,1,name,size=10)
    cell(ws,r,2,sifs(CR,(NM,f"$A{r}"),(FL,q("REVENUE")),(LN,q("Аренда"))),align='right',fmt=NUM)
    cell(ws,r,3,sifs(CR,(NM,f"$A{r}"),(FL,q("REVENUE")),(LN,q("ТЭО"))),align='right',fmt=NUM)
    cell(ws,r,4,f"=COUNTIFS({NM},$A{r},{FL},{q('REVENUE')})",align='center')
    cell(ws,r,5,f"=B{r}+C{r}",align='right',fmt=NUM)
    cell(ws,r,6,f"=IFERROR(E{r}/$E${TOTv},0)",align='right',fmt=PCT)
r=TOTv
cell(ws,r,1,"ИТОГО",bold=True,fill=LGREY)
cell(ws,r,2,sifs(CR,(FL,q("REVENUE")),(LN,q("Аренда"))),bold=True,align='right',fmt=NUM,fill=LGREY)
cell(ws,r,3,sifs(CR,(FL,q("REVENUE")),(LN,q("ТЭО"))),bold=True,align='right',fmt=NUM,fill=LGREY)
cell(ws,r,4,f"=COUNTIFS({FL},{q('REVENUE')})",bold=True,align='center',fill=LGREY)
cell(ws,r,5,f"=B{r}+C{r}",bold=True,align='right',fmt=NUM,fill=LGREY)
cell(ws,r,6,"=1",bold=True,align='right',fmt=PCT,fill=LGREY)
ws.column_dimensions['A'].width=46
for c in 'BCE': ws.column_dimensions[c].width=16
ws.column_dimensions['D'].width=10; ws.column_dimensions['F'].width=9
ws.freeze_panes="A5"

# ============================================================
# 6. РАСХОДЫ
# ============================================================
ws=wb.create_sheet("Расходы"); ws.sheet_view.showGridLines=False
cell(ws,1,1,"РАСХОДЫ ПО КАТЕГОРИЯМ И ПОСТАВЩИКАМ",bold=True,size=14,fontcolor=NAVY,bd=False)
cell(ws,2,1,f"{PERIOD}  ·  ₸",size=10,italic=True,bd=False)
r=4; banner(ws,r,2,"По категориям",fill=BLUE); r+=1
for c,t in [(1,"Категория"),(2,"Сумма, ₸")]: hdr(ws,r,c,t)
catlist=[("COGS",c) for c in cogs_cats]+[("OPEX",c) for c in opex_cats]
start=r+1
for fl,cat in catlist:
    r+=1; cell(ws,r,1,("COGS · " if fl=="COGS" else "OPEX · ")+cat,size=10)
    cell(ws,r,2,sifs(DB,(FL,q(fl)),(CT,q(cat))),align='right',fmt=NUM)
r+=1; cell(ws,r,1,"Налоги и соц. платежи",size=10); cell(ws,r,2,sifs(DB,(FL,q("TAX"))),align='right',fmt=NUM)
r+=1; cell(ws,r,1,"Проценты и пени по кредитам",size=10); cell(ws,r,2,sifs(DB,(FL,q("FIN_COST"))),align='right',fmt=NUM)
last_c=r
r+=1; cell(ws,r,1,"ИТОГО (без займов/депозитов/переводов)",bold=True,fill=LGREY)
cell(ws,r,2,f"=SUM(B{start}:B{last_c})",bold=True,align='right',fmt=NUM,fill=LGREY)
r+=2; banner(ws,r,2,"Топ-20 поставщиков (COGS + OPEX)",fill=BLUE); r+=1
for c,t in [(1,"Поставщик"),(2,"Сумма, ₸")]: hdr(ws,r,c,t)
for name in suppliers:
    r+=1; cell(ws,r,1,name,size=10)
    cell(ws,r,2,f"={sifs(DB,(NM,f'$A{r}'),(FL,q('COGS')))[1:]}+{sifs(DB,(NM,f'$A{r}'),(FL,q('OPEX')))[1:]}",align='right',fmt=NUM)
ws.column_dimensions['A'].width=54; ws.column_dimensions['B'].width=18

# order: Резюме first
wb.move_sheet("Резюме", -(wb.sheetnames.index("Резюме")))
wb.active=0
out='/home/user/board-presentation/mgmt-report/WayStar_Управленческая_отчетность_15.06-15.07.2026.xlsx'
wb.save(out)
print("SAVED:",out,"| sheets:",wb.sheetnames)
