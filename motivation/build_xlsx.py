# -*- coding: utf-8 -*-
"""Waystar Group — 75/25 KPI premium calculator (.xlsx).
Model: salary = 75% fixed + 25% variable. Variable = (25%*salary)*KPI%*company_mod*gate, cap 150%.
Front-line roles carry a personal/team gross-or-net-profit KPI; back-office carries a company-profit KPI.
Sheet 'Калькулятор' — pay per employee; Sheet 'Детализация KPI' — per-KPI achievement roll-up."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY='13294B'; STEEL='2F6690'; COPPER='C4772F'; LIGHT='F1F5F9'; WHITE='FFFFFF'; GREEN='2E7D46'; YELLOW='FFF3C4'; PROF='E8F3EC'
thin=Side(style='thin',color='C9D3DE'); border=Border(left=thin,right=thin,top=thin,bottom=thin)
def fill(c): return PatternFill('solid',fgColor=c)
HDR=Font(name='Calibri',bold=True,color=WHITE,size=10)
BOLD=Font(name='Calibri',bold=True,size=10,color='2A3746')
REG=Font(name='Calibri',size=10,color='2A3746')
CTR=Alignment(horizontal='center',vertical='center',wrap_text=True)
LEFT=Alignment(horizontal='left',vertical='center',wrap_text=True)
RIGHT=Alignment(horizontal='right',vertical='center')

# idx, name, role, type(front/back), oklad_gross, [(kpi, weight, is_profit)]
EMP=[
 (1,'Аманбекова Айгуль','HR (директор по персоналу)','back',1246243,
   [('Чистая прибыль компании (общий)',0.25,1),('Удержание персонала / текучесть',0.25,0),
    ('Закрытие ключевых вакансий в срок',0.20,0),('eNPS / вовлечённость',0.15,0),('Запуск KPI / ФОТ-бюджет',0.15,0)]),
 (2,'Аубакиров Алимбек','CFO','front',4347889,
   [('Чистая прибыль компании vs план',0.35,1),('Ликвидность / cash runway',0.20,0),('DSO — дебиторка',0.15,0),
    ('Точность бюджета',0.15,0),('Привлечение финансирования',0.15,0)]),
 (3,'Байльденова Шолпан','HR менеджер','back',614931,
   [('Чистая прибыль компании (общий)',0.25,1),('Скорость подбора',0.20,0),('Кадровый документооборот',0.20,0),
    ('Точность payroll',0.20,0),('Онбординг / испыт. срок',0.15,0)]),
 (4,'Данияр Алим','CEO','front',1246243,
   [('Чистая прибыль компании vs план',0.40,1),('Валовая прибыль / маржа',0.20,1),('Выручка vs план',0.20,0),
    ('Бизнес-план / OKR',0.15,0),('Кросс-функц. / eNPS',0.05,0)]),
 (5,'Доолоткулов Данияр','Начальник отдела логистики','front',1246243,
   [('Валовая прибыль отдела логистики vs план',0.45,1),('On-time delivery',0.20,0),
    ('Стоимость доставки на тонну',0.20,0),('Потери / претензии',0.15,0)]),
 (6,'Жунусова Камила','Менеджер по логистике','front',230896,
   [('Валовая маржа по своим рейсам (≥K×ФОТ)',0.50,1),('Обработанные заявки/рейсы',0.20,0),
    ('On-time %',0.15,0),('Точность документов / претензии',0.15,0)]),
 (7,'Сейтқадыр Дарын','Менеджер по логистике','front',236143,
   [('Валовая маржа по своим рейсам (≥K×ФОТ)',0.50,1),('Обработанные заявки/рейсы',0.20,0),
    ('On-time %',0.15,0),('Точность документов / претензии',0.15,0)]),
 (8,'Серикбаев Адиль','Президент','front',6000139,
   [('Чистая прибыль компании vs план',0.50,1),('Выручка vs план',0.15,0),('Стратег. инициативы',0.20,0),
    ('Финансирование / ДДС',0.10,0),('Удержание топ-команды',0.05,0)]),
 (9,'Таджимурадова Асем','Менеджер по закрытию','front',378788,
   [('Валовая прибыль по закрытым сделкам vs план',0.50,1),('Собираемость дебиторки',0.25,0),
    ('Срок закрытия документов',0.15,0),('Точность / 0 расхождений',0.10,0)]),
 (10,'Ташметов Сакен','Юрист','back',1246243,
   [('Чистая прибыль компании (общий)',0.25,1),('Взысканная дебиторка / споры',0.25,0),('Договоры в SLA',0.20,0),
    ('Юр. риски: штрафы',0.15,0),('Поддержка сделок (DD)',0.15,0)]),
 (11,'Унбетпаев Ернар','Операционный директор','front',1873611,
   [('Валовая прибыль по операциям vs план',0.45,1),('OTIF',0.20,0),
    ('Удельная себестоимость операций',0.20,0),('HSE / качество',0.15,0)]),
 (12,'Белова Полина','Менеджер по логистике (junior)','front',173011,
   [('Валовая маржа по своим рейсам (≥K×ФОТ)',0.50,1),('Обработанные заявки/рейсы',0.20,0),
    ('On-time %',0.15,0),('Точность документов / претензии',0.15,0)]),
 (13,'Сибагатов Хамардин','Менеджер по логистике (senior)','front',993718,
   [('Валовая маржа по своим рейсам (≥K×ФОТ)',0.50,1),('Обработанные заявки/рейсы',0.20,0),
    ('On-time %',0.15,0),('Точность документов / претензии',0.15,0)]),
]
FIX=0.75; VAR=0.25; CAP=1.5

wb=Workbook()

# ---------- Детализация KPI ----------
det=wb.active; det.title='Детализация KPI'; det.sheet_view.showGridLines=False
det['A1']='Waystar Group · Детализация KPI (модель 75/25)'
det['A1'].font=Font(name='Calibri',bold=True,size=14,color=NAVY)
det['A2']='Введите фактическое достижение по каждому KPI (жёлтые ячейки, 0–150%). Зелёные строки — профит-цель (валовая/чистая прибыль). '\
          'Итоговый %KPI = Σ(вес×достижение) переносится в лист «Калькулятор».'
det['A2'].font=Font(name='Calibri',size=9,italic=True,color='6B7A8D'); det.merge_cells('A2:F2')
dv=DataValidation(type='decimal',operator='between',formula1=0,formula2=1.5,allow_blank=True)
dv.error='Достижение 0–150%'; dv.errorTitle='Диапазон'; det.add_data_validation(dv)
for col,w in {'A':4,'B':44,'C':10,'D':16,'E':16,'F':4}.items(): det.column_dimensions[col].width=w

subtotal_ref={}
r=4
for (idx,name,role,typ,oklad,kpis) in EMP:
    tag='ФРОНТ' if typ=='front' else 'БЭК'
    det.cell(r,2,'%d. %s — %s  [%s]'%(idx,name,role,tag))
    for col in range(1,7):
        det.cell(r,col).fill=fill(STEEL); det.cell(r,col).border=border
        det.cell(r,col).font=Font(name='Calibri',bold=True,size=10,color=WHITE)
    r+=1
    det.cell(r,2,'KPI').font=HDR; det.cell(r,3,'Вес').font=HDR
    det.cell(r,4,'Достижение, %').font=HDR; det.cell(r,5,'Вклад').font=HDR
    for col in range(1,7):
        det.cell(r,col).fill=fill(NAVY); det.cell(r,col).border=border; det.cell(r,col).alignment=CTR
    r+=1; first=r
    for (kpi,w,isp) in kpis:
        det.cell(r,2,kpi).font=(Font(name='Calibri',bold=True,size=10,color=GREEN) if isp else REG); det.cell(r,2).alignment=LEFT
        wc=det.cell(r,3,w); wc.number_format='0%'; wc.font=REG; wc.alignment=CTR
        ac=det.cell(r,4,1.0); ac.number_format='0%'; ac.font=Font(name='Calibri',size=10,color='1F4E79',bold=True)
        ac.fill=fill(YELLOW); ac.alignment=CTR; dv.add(ac)
        cc=det.cell(r,5,'=C%d*D%d'%(r,r)); cc.number_format='0%'; cc.font=REG; cc.alignment=CTR
        for col in range(1,7):
            det.cell(r,col).border=border
            if isp and col in (2,3,5): det.cell(r,col).fill=fill(PROF)
        r+=1
    last=r-1
    det.cell(r,2,'Итоговый %KPI').font=BOLD
    wsum=det.cell(r,3,'=SUM(C%d:C%d)'%(first,last)); wsum.number_format='0%'; wsum.font=BOLD; wsum.alignment=CTR
    tcell=det.cell(r,5,'=SUM(E%d:E%d)'%(first,last)); tcell.number_format='0%'
    tcell.font=Font(name='Calibri',bold=True,size=10,color=GREEN); tcell.alignment=CTR
    for col in range(1,7): det.cell(r,col).fill=fill(LIGHT); det.cell(r,col).border=border
    subtotal_ref[idx]=r
    r+=2
det.freeze_panes='A4'

# ---------- Калькулятор ----------
calc=wb.create_sheet('Калькулятор',0); calc.sheet_view.showGridLines=False
calc['A1']='Waystar Group · Калькулятор оплаты 75/25 по KPI'
calc['A1'].font=Font(name='Calibri',bold=True,size=15,color=NAVY)
calc['A2']='Итого = 75%×Оклад + (25%×Оклад)×%KPI×Модификатор компании×Гейт   (переменная — кап 150%)'
calc['A2'].font=Font(name='Calibri',size=10,italic=True,color=COPPER)
calc['A4']='Модификатор компании (по чистой прибыли, 0,7–1,3):'; calc['A4'].font=BOLD
mod=calc['F4']; mod.value=1.0; mod.number_format='0.00'; mod.font=Font(bold=True,color='1F4E79',size=11)
mod.fill=fill(YELLOW); mod.alignment=CTR; mod.border=border
calc['G4']='← 100% плана прибыли → 1,00'; calc['G4'].font=Font(size=9,italic=True,color='6B7A8D')
dvm=DataValidation(type='decimal',operator='between',formula1=0,formula2=1.3); calc.add_data_validation(dvm); dvm.add(mod)

heads=['№','Сотрудник','Должность','Тип','Оклад (100%), ₸','Фикс 75%, ₸','Перем. target 25%, ₸',
       '%KPI','Гейт','Перем. к выплате, ₸','ИТОГО к выплате, ₸','Δ к текущему, ₸']
hr=6
for j,txt in enumerate(heads,1):
    c=calc.cell(hr,j,txt); c.font=HDR; c.fill=fill(NAVY); c.border=border; c.alignment=CTR
widths=[4,24,27,7,14,13,15,9,7,15,16,14]
for j,w in enumerate(widths,1): calc.column_dimensions[get_column_letter(j)].width=w
dvg=DataValidation(type='list',formula1='"1,0.5,0"'); calc.add_data_validation(dvg)

row=hr+1; first_data=row
for (idx,name,role,typ,oklad,kpis) in EMP:
    calc.cell(row,1,idx).font=REG; calc.cell(row,1).alignment=CTR
    calc.cell(row,2,name).font=REG; calc.cell(row,2).alignment=LEFT
    calc.cell(row,3,role).font=REG; calc.cell(row,3).alignment=LEFT
    tg=calc.cell(row,4,'ФРОНТ' if typ=='front' else 'БЭК'); tg.alignment=CTR
    tg.font=Font(name='Calibri',bold=True,size=9,color=(GREEN if typ=='front' else STEEL))
    oc=calc.cell(row,5,oklad); oc.number_format='#,##0'; oc.font=REG; oc.alignment=RIGHT
    fx=calc.cell(row,6,'=E%d*%s'%(row,FIX)); fx.number_format='#,##0'; fx.font=REG; fx.alignment=RIGHT
    vt=calc.cell(row,7,'=E%d*%s'%(row,VAR)); vt.number_format='#,##0'; vt.font=REG; vt.alignment=RIGHT
    kc=calc.cell(row,8,"='Детализация KPI'!E%d"%subtotal_ref[idx]); kc.number_format='0%'
    kc.font=Font(name='Calibri',size=10,color=GREEN,bold=True); kc.alignment=CTR
    gc=calc.cell(row,9,1); gc.number_format='0.0'; gc.font=Font(color='1F4E79',bold=True); gc.alignment=CTR
    gc.fill=fill(YELLOW); dvg.add(gc)
    # variable earned = target * MIN(cap, kpi*mod*gate)
    ve=calc.cell(row,10,'=G%d*MIN(%s,H%d*$F$4*I%d)'%(row,CAP,row,row)); ve.number_format='#,##0'; ve.font=BOLD; ve.alignment=RIGHT
    tot=calc.cell(row,11,'=F%d+J%d'%(row,row)); tot.number_format='#,##0'; tot.font=BOLD; tot.alignment=RIGHT
    dl=calc.cell(row,12,'=K%d-E%d'%(row,row)); dl.number_format='#,##0;[Red]-#,##0'; dl.font=REG; dl.alignment=RIGHT
    zb=LIGHT if (row-first_data)%2==0 else WHITE
    for j in range(1,13):
        calc.cell(row,j).border=border
        if j not in (8,9): calc.cell(row,j).fill=fill(zb)
    row+=1
last_data=row-1
calc.cell(row,2,'ИТОГО').font=Font(bold=True,color=WHITE,size=10)
for col,f in [(5,'E'),(6,'F'),(7,'G'),(10,'J'),(11,'K'),(12,'L')]:
    c=calc.cell(row,col,'=SUM(%s%d:%s%d)'%(f,first_data,f,last_data)); c.number_format='#,##0'
for j in range(1,13):
    calc.cell(row,j).fill=fill(STEEL); calc.cell(row,j).border=border
    calc.cell(row,j).font=Font(bold=True,color=WHITE,size=10)
    calc.cell(row,j).alignment=RIGHT if j in (5,6,7,10,11,12) else CTR
trow=row
row+=2
calc.cell(row,2,'ФОТ (оклады) в год, ₸:').font=BOLD
calc.cell(row,5,'=E%d*12'%trow).number_format='#,##0'; calc.cell(row,5).font=Font(bold=True,color=NAVY); calc.cell(row,5).alignment=RIGHT
calc.cell(row,10,'Под риском (25%) в год:').font=BOLD
calc.cell(row,11,'=G%d*12'%trow).number_format='#,##0'; calc.cell(row,11).font=Font(bold=True,color=COPPER); calc.cell(row,11).alignment=RIGHT
row+=1
calc.cell(row,2,'Итого к выплате в год, ₸:').font=BOLD
calc.cell(row,11,'=K%d*12'%trow).number_format='#,##0'; calc.cell(row,11).font=Font(bold=True,color=GREEN,size=11); calc.cell(row,11).alignment=RIGHT
row+=2
for note in [
 'Жёлтые ячейки — вводимые: Модификатор компании (F4), Гейт по строке, достижения KPI на листе «Детализация KPI».',
 'Тип ФРОНТ — есть личная/командная профит-цель (валовая/чистая прибыль); БЭК — доля переменной на чистой прибыли компании.',
 'Переменная = (25%×оклад) × %KPI × модификатор × гейт, ограничена сверху 150% target. При 100% и модификаторе 1,0 «Итого» = текущему окладу.',
 'Оклад — начисленный (база деления 75/25). Пороги KPI и коэффициент K профит-целей калибруются по P&L.']:
    calc.cell(row,2,'• '+note).font=Font(size=9,italic=True,color='6B7A8D')
    calc.merge_cells(start_row=row,start_column=2,end_row=row,end_column=12); row+=1

calc.freeze_panes='A7'
wb.save('motivation/Waystar_KPI_premium_calculator.xlsx')
print('saved motivation/Waystar_KPI_premium_calculator.xlsx')
